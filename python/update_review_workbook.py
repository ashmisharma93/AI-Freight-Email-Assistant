
import csv
from pathlib import Path
from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parent.parent
CSV_FILE = PROJECT_DIR / "output" / "gmail_test.csv"
EXCEL_FILE = PROJECT_DIR / "output" / "Freight_Email_Review.xlsx"

REVIEW_COLUMN = "review_status"
ID_COLUMN = "email_id"


def update_workbook():
    if not CSV_FILE.exists():
        raise FileNotFoundError(f"CSV file not found: {CSV_FILE}")

    if not EXCEL_FILE.exists():
        raise FileNotFoundError(f"Excel workbook not found: {EXCEL_FILE}")

    # Read incoming email records from the CSV.
    with CSV_FILE.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        csv_headers = reader.fieldnames

        if not csv_headers or ID_COLUMN not in csv_headers:
            raise ValueError("The CSV must contain an email_id column.")

        incoming_records = list(reader)

    # Load the existing review workbook without replacing it.
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    headers = [
        sheet.cell(row=1, column=column).value
        for column in range(1, sheet.max_column + 1)
    ]

    if ID_COLUMN not in headers or REVIEW_COLUMN not in headers:
        workbook.close()
        raise ValueError(
            "The Excel sheet must contain email_id and review_status columns."
        )

    # Map existing emails to their Excel row numbers.
    id_column_number = headers.index(ID_COLUMN) + 1
    review_column_number = headers.index(REVIEW_COLUMN) + 1

    existing_rows = {}
    for row_number in range(2, sheet.max_row + 1):
        email_id = sheet.cell(row=row_number, column=id_column_number).value
        if email_id:
            existing_rows[str(email_id)] = row_number

    added = 0
    skipped = 0

    # Append only new emails; never overwrite existing review decisions.
    for record in incoming_records:
        email_id = (record.get(ID_COLUMN) or "").strip()

        if not email_id or email_id in existing_rows:
            skipped += 1
            continue

        new_row = sheet.max_row + 1

        for column_number, header in enumerate(headers, start=1):
            if header in csv_headers:
                value = record.get(header, "")

                if header == REVIEW_COLUMN:
                    value = "Pending Review"

                sheet.cell(
                    row=new_row,
                    column=column_number
                ).value = value

        # Copy the previous row's formatting to the new row.
        if new_row > 2:
            for column_number in range(1, len(headers) + 1):
                source = sheet.cell(
                    row=new_row - 1, column=column_number
                )
                target = sheet.cell(
                    row=new_row, column=column_number
                )

                if source.has_style:
                    from copy import copy
                    target._style = copy(source._style)

                if source.number_format:
                    target.number_format = source.number_format

        existing_rows[email_id] = new_row
        added += 1

    # Save changes back into the existing workbook.
    for table in sheet.tables.values():
        table.ref = (
            f"A1:{sheet.cell(row=1, column=sheet.max_column).column_letter}"
            f"{sheet.max_row}"
        )

    workbook.save(EXCEL_FILE)
    workbook.close()

    print(f"Update complete. New emails added: {added}")
    print(f"Existing or invalid email IDs skipped: {skipped}")
    print("Existing review statuses have been preserved.")


if __name__ == "__main__":
    update_workbook()
